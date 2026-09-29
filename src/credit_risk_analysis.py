"""Retail Credit Risk & Loan Demand Analytics (simulated data).
Run from the repo root: python src/credit_risk_analysis.py  -> rebuilds data, SQL, Excel and dashboard.
Needs: pandas, numpy, openpyxl (sqlite3 is built into Python)."""
import os, json, sqlite3
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MONTHS = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']

# ---------- 1. SIMULATED RAW DATA (with deliberate data-quality issues) ----------
def make_data():
    rng = np.random.default_rng(42); N = 100_000
    days = pd.date_range('2023-01-01', '2025-12-31')
    seas = np.array([.9,.85,1.1,.9,.85,.8,.85,.95,1.05,1.35,1.3,1.15])
    w = seas[np.asarray(days.month) - 1] * (1 + 0.2 * np.asarray((days - days[0]).days) / 365)
    d = days[rng.choice(len(days), N, p=w / w.sum())]
    yr = np.asarray(d.year) - 2023
    RA = {'North': 0.0, 'South': -0.1, 'East': 0.2, 'West': -0.05, 'Central': 0.15}
    reg = rng.choice(list(RA), N, p=[.24,.28,.14,.24,.10])
    radj = pd.Series(reg).map(RA).values
    score = np.clip(rng.normal(690, 75, N), 300, 850).round().astype(int)
    inc = np.exp(rng.normal(13.2, .5, N)); loan = np.round(inc * rng.uniform(.2, 1.5, N), -3); inc = inc.round(-3)
    dti = np.clip(rng.normal(.33, .11, N), .02, .85).round(3)
    util = np.clip(rng.beta(2, 3.2, N), 0, 1).round(3)
    sig = lambda z: 1 / (1 + np.exp(-z))
    appr = (rng.random(N) < sig((score - 640) / 45 - (dti - .35) * 5)).astype(int)
    pdef = sig(-2.7 + (650 - score) / 70 + (dti - .33) * 5 + (util - .4) * 1.5 + radj)
    dflt = ((rng.random(N) < pdef) & (appr == 1)).astype(int)
    fraud = (rng.random(N) < .006 + .012 * (util > .8) + .004 * (loan / inc > 1.2)).astype(int)
    df = pd.DataFrame(dict(
        loan_id=['LN%d' % (100000 + i) for i in range(N)], application_date=d.strftime('%Y-%m-%d'),
        region=reg, purpose=rng.choice(['Debt Consolidation','Personal','Vehicle','Home Improvement','Business','Education'], N, p=[.28,.18,.15,.15,.12,.12]),
        credit_score=score, annual_income=inc, loan_amount=loan, dti=dti, credit_utilization=util,
        approved=appr, defaulted=dflt, fraud_flag=fraud,
        processing_days=np.clip(rng.gamma(2.5, 1.7 - .15 * yr), .5, 15).round(1)))
    m = lambda p: rng.random(N) < p
    df.loc[m(.03), 'annual_income'] = np.nan          # missing income
    df.loc[m(.02), 'dti'] = np.nan                    # missing DTI
    df.loc[m(.003), 'annual_income'] = df['annual_income'] * 60   # data-entry outliers
    df.loc[m(.001), 'credit_score'] = 0               # invalid scores
    return pd.concat([df, df.sample(500, random_state=1)], ignore_index=True)   # duplicates

# ---------- 2. SQL ----------
CLEAN = """DROP TABLE IF EXISTS loans_clean;
CREATE TABLE loans_clean AS
WITH dedup AS (            -- remove duplicate loan_ids
  SELECT * FROM loans_raw WHERE rowid IN (SELECT MIN(rowid) FROM loans_raw GROUP BY loan_id)
), valid AS (              -- drop impossible credit scores
  SELECT * FROM dedup WHERE credit_score BETWEEN 300 AND 850
), stats AS (              -- regional averages used to fill gaps
  SELECT region, AVG(dti) AS avg_dti,
         AVG(CASE WHEN annual_income <= 10000000 THEN annual_income END) AS avg_inc
  FROM valid GROUP BY region
), fixed AS (              -- impute missing income/DTI, treat income > 1 crore as outlier
  SELECT v.loan_id, v.application_date, v.region, v.purpose, v.credit_score,
         CAST(ROUND(CASE WHEN v.annual_income IS NULL OR v.annual_income > 10000000
                         THEN s.avg_inc ELSE v.annual_income END) AS INTEGER) AS annual_income,
         v.loan_amount, ROUND(COALESCE(v.dti, s.avg_dti), 3) AS dti, v.credit_utilization,
         v.approved, v.defaulted, v.fraud_flag, v.processing_days
  FROM valid v JOIN stats s ON s.region = v.region
)
SELECT *, strftime('%Y-%m', application_date) AS app_month,
  CASE WHEN credit_score >= 760 THEN 'A' WHEN credit_score >= 720 THEN 'B'
       WHEN credit_score >= 680 THEN 'C' WHEN credit_score >= 620 THEN 'D' ELSE 'E' END AS grade,
  CASE WHEN dti < 0.2 THEN '<20%' WHEN dti < 0.3 THEN '20-30%' WHEN dti < 0.4 THEN '30-40%' ELSE '40%+' END AS dti_band,
  CASE WHEN credit_utilization < 0.3 THEN '<30%' WHEN credit_utilization < 0.6 THEN '30-60%'
       WHEN credit_utilization < 0.8 THEN '60-80%' ELSE '80%+' END AS util_band,
  CASE WHEN annual_income < 400000 THEN 'Low (<4L)' WHEN annual_income < 800000 THEN 'Mid (4-8L)'
       WHEN annual_income < 1500000 THEN 'Upper-mid (8-15L)' ELSE 'High (15L+)' END AS income_bracket
FROM fixed;"""

def seg(col, order=None, extra='', having=''):
    return (f"SELECT {col} AS segment, COUNT(*) AS applications, SUM(approved) AS approved, SUM(defaulted) AS defaults{extra}\n"
            f"FROM loans_clean GROUP BY {col} {having} ORDER BY {order or col};")

Q = {
 'kpi': """SELECT COUNT(*) AS applications, SUM(approved) AS approved, SUM(defaulted) AS defaults,
  ROUND(SUM(CASE WHEN approved=1 THEN loan_amount END)/1e7,0) AS disbursed_cr,
  ROUND(SUM(CASE WHEN defaulted=1 THEN loan_amount*0.6 END)/1e7,0) AS est_loss_cr,
  ROUND(100.0*SUM(fraud_flag)/COUNT(*),2) AS fraud_pct,
  ROUND(100.0*SUM(processing_days>5)/COUNT(*),1) AS sla_breach_pct FROM loans_clean;""",
 'grade': seg('grade'),
 'dti': seg('dti_band', 'MIN(dti)'),
 'util': seg('util_band', 'MIN(credit_utilization)'),
 'income': seg('income_bracket', 'MIN(annual_income)'),
 'purpose': seg('purpose'),
 'region': seg('region'),
 'monthly': seg('app_month'),
 'season': seg("substr(app_month,6,2)"),
 'region_ops': """SELECT region AS segment, COUNT(*) AS applications,
  ROUND(100.0*SUM(approved)/COUNT(*),1) AS approval_pct,
  ROUND(100.0*SUM(defaulted)/SUM(approved),2) AS default_pct,
  ROUND(AVG(processing_days),1) AS avg_days, ROUND(100.0*SUM(processing_days>5)/COUNT(*),1) AS sla_breach_pct,
  ROUND(100.0*SUM(fraud_flag)/COUNT(*),2) AS fraud_pct FROM loans_clean GROUP BY region ORDER BY default_pct DESC;""",
 'sla_year': """SELECT substr(app_month,1,4) AS segment, ROUND(100.0*SUM(processing_days>5)/COUNT(*),1) AS sla_breach_pct,
  ROUND(AVG(processing_days),2) AS avg_days FROM loans_clean GROUP BY 1 ORDER BY 1;""",
 'risk_pockets': """SELECT grade || ' grade, DTI ' || dti_band AS segment, COUNT(*) AS applications, SUM(approved) AS approved,
  SUM(defaulted) AS defaults FROM loans_clean GROUP BY grade, dti_band HAVING SUM(approved) >= 150
  ORDER BY 1.0*SUM(defaulted)/SUM(approved) DESC LIMIT 6;""",
}

def rates(df):
    df = df.copy(); df['approval_pct'] = 100 * df.approved / df.applications
    df['default_pct'] = 100 * df.defaults / df.approved.replace(0, np.nan); return df

# ---------- 3. RUN PIPELINE ----------
raw = make_data()
con = sqlite3.connect(':memory:')
raw.to_sql('loans_raw', con, index=False)
con.executescript(CLEAN)
R = {k: pd.read_sql(v, con) for k, v in Q.items()}
clean = pd.read_sql('SELECT * FROM loans_clean', con)
clean.to_csv(f'{ROOT}/data/loans_clean.csv', index=False)          # import this into Power BI
with open(f'{ROOT}/sql/queries.sql', 'w') as f:
    f.write('-- SQLite. Load loans_clean.csv/raw data into table loans_raw first.\n-- STEP 1: CLEANING + SEGMENTATION\n' + CLEAN)
    for k, v in Q.items(): f.write(f'\n\n-- ANALYSIS: {k}\n{v}')
for k in ['grade','dti','util','income','purpose','region','monthly','season','risk_pockets']: R[k] = rates(R[k])
K = R['kpi'].iloc[0]
n_raw, n_clean = len(raw), len(clean)
dups = 500; bad = int((raw.drop_duplicates('loan_id').credit_score == 0).sum())
n_inc = int((raw.drop_duplicates('loan_id').annual_income.isna() | (raw.drop_duplicates('loan_id').annual_income > 1e7)).sum())
n_dti = int(raw.drop_duplicates('loan_id').dti.isna().sum())

# ---------- 4. FINDINGS ----------
g, dt, ut, rg, se, sl = R['grade'], R['dti'], R['util'], R['region_ops'], R['season'], R['sla_year']
mon = R['monthly']
pk = se.sort_values('applications', ascending=False).segment.head(3).astype(int).map(lambda m: MONTHS[m-1]).tolist()
lo = MONTHS[int(se.sort_values('applications').segment.iloc[0]) - 1]
F = [
 f"Credit grade is the strongest risk driver: Grade E defaults at {g.default_pct.iloc[-1]:.1f}% vs {g.default_pct.iloc[0]:.1f}% for Grade A ({g.default_pct.iloc[-1]/g.default_pct.iloc[0]:.1f}x).",
 f"Debt-to-income matters: borrowers at DTI 40%+ default at {dt.default_pct.iloc[-1]:.1f}% vs {dt.default_pct.iloc[0]:.1f}% below 20%.",
 f"Credit utilisation above 80% is a warning sign: {ut.default_pct.iloc[-1]:.1f}% default vs {ut.default_pct.iloc[0]:.1f}% under 30%.",
 f"Highest-risk region is {rg.segment.iloc[0]} ({rg.default_pct.iloc[0]:.1f}% default); lowest is {rg.segment.iloc[-1]} ({rg.default_pct.iloc[-1]:.1f}%).",
 f"Demand peaks in {', '.join(pk)} (festive/year-end season) and is lowest in {lo}: plan staffing and credit-check capacity ahead of Q4.",
 f"SLA breaches (decision > 5 days) fell from {sl.sla_breach_pct.iloc[0]:.1f}% in {sl.segment.iloc[0]} to {sl.sla_breach_pct.iloc[-1]:.1f}% in {sl.segment.iloc[-1]}, but remain material.",
]
REC = ["Tighten approval rules or pricing for Grade D/E applicants with DTI above 40% (see risk pockets).",
       "Add a utilisation cap (>80%) as a secondary review trigger.",
       "Increase processing staff before October; use monthly demand as a staffing forecast."]

# ---------- 5. EXCEL ----------
wb = Workbook(); wd = wb.active; wd.title = 'Dashboard'; wa = wb.create_sheet('Analysis'); wc = wb.create_sheet('Clean_Data_10K'); wn = wb.create_sheet('About')
FN = 'Arial'; NAVY = '1F3864'
pos = {}; row = 1
def put(name, title, df):
    global row
    wa.cell(row, 1, title).font = Font(name=FN, bold=True, size=12, color=NAVY)
    hdr = row + 1
    for j, c in enumerate(['segment','applications','approved','defaults','approval_rate','default_rate'], 1):
        x = wa.cell(hdr, j, c); x.font = Font(name=FN, bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=NAVY)
    for i, r in enumerate(df[['segment','applications','approved','defaults']].itertuples(index=False), hdr + 1):
        for j, v in enumerate(r, 1):
            wa.cell(i, j, v.item() if hasattr(v, 'item') else v).font = Font(name=FN)
        wa.cell(i, 5, f'=IF(B{i}=0,0,C{i}/B{i})').number_format = '0.0%'
        wa.cell(i, 6, f'=IF(C{i}=0,0,D{i}/C{i})').number_format = '0.0%'
        for j in (5, 6): wa.cell(i, j).font = Font(name=FN)
    pos[name] = (hdr, hdr + len(df)); row = hdr + len(df) + 3
put('grade', 'Default rate by credit grade (A = best)', R['grade'])
put('dti', 'Default rate by debt-to-income band', R['dti'])
put('util', 'Default rate by credit utilisation band', R['util'])
put('income', 'Default rate by income bracket (INR)', R['income'])
put('purpose', 'Default rate by loan purpose', R['purpose'])
put('region', 'Default rate by region', R['region'])
put('monthly', 'Monthly loan applications (demand trend)', R['monthly'])
put('risk_pockets', 'Top 6 highest-risk segments (min 150 approved loans)', R['risk_pockets'])
wa.column_dimensions['A'].width = 28
for c in 'BCDEF': wa.column_dimensions[c].width = 15

def chart(kind, name, valcol, title, anchor, ytitle):
    h, l = pos[name]; ch = BarChart() if kind == 'bar' else LineChart()
    ch.title = title; ch.height = 7.5; ch.width = 16; ch.legend = None
    ch.add_data(Reference(wa, min_col=valcol, min_row=h, max_row=l), titles_from_data=True)
    ch.set_categories(Reference(wa, min_col=1, min_row=h + 1, max_row=l))
    ch.y_axis.title = ytitle; ch.x_axis.delete = False; ch.y_axis.delete = False
    if valcol == 6: ch.y_axis.number_format = '0%'
    wd.add_chart(ch, anchor)

wd['A1'] = 'Retail Credit Risk & Loan Demand Dashboard (simulated data)'; wd['A1'].font = Font(name=FN, bold=True, size=16, color=NAVY)
gh, gl = pos['grade']
kp = [('Applications', f'=SUM(Analysis!B{gh+1}:B{gl})', '#,##0'), ('Approved loans', f'=SUM(Analysis!C{gh+1}:C{gl})', '#,##0'),
      ('Approval rate', '=B4/A4', '0.0%'), ('Defaults', f'=SUM(Analysis!D{gh+1}:D{gl})', '#,##0'), ('Default rate', '=D4/B4', '0.0%')]
for j, (lab, fm, nf) in enumerate(kp, 1):
    a = wd.cell(3, j, lab); a.font = Font(name=FN, bold=True, color='FFFFFF'); a.fill = PatternFill('solid', fgColor=NAVY); a.alignment = Alignment(horizontal='center')
    b = wd.cell(4, j, fm); b.number_format = nf; b.font = Font(name=FN, bold=True, size=14); b.alignment = Alignment(horizontal='center')
    wd.column_dimensions[chr(64 + j)].width = 18
chart('bar', 'grade', 6, 'Default rate by credit grade', 'A6', 'Default rate')
chart('bar', 'dti', 6, 'Default rate by DTI band', 'H6', 'Default rate')
chart('line', 'monthly', 2, 'Monthly loan applications', 'A22', 'Applications')
chart('bar', 'region', 6, 'Default rate by region', 'H22', 'Default rate')
wd['A38'] = 'Key findings'; wd['A38'].font = Font(name=FN, bold=True, size=12, color=NAVY)
for i, t in enumerate(F, 39): wd.cell(i, 1, '- ' + t).font = Font(name=FN)

wc.append(list(clean.columns))
for r in clean.sample(10000, random_state=1).itertuples(index=False): wc.append(list(r))
for c in wc[1]: c.font = Font(name=FN, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY)
wc.freeze_panes = 'A2'
for t in ['ABOUT THIS WORKBOOK', 'Data is SIMULATED (seeded random generator) to mimic a retail lending portfolio. Not real customer data.',
          f'Full analysis covers {n_clean:,} cleaned applications; Clean_Data_10K is a random 10,000-row sample.',
          'Analysis tab: counts come from SQL (queries.sql); approval and default rates are live Excel formulas.',
          'Dashboard tab: KPI cards and charts read from the Analysis tab.', 'Assumption: loss given default = 60% of loan amount; SLA breach = decision took more than 5 days.',
          'Currency: INR. 1 crore (Cr) = 10 million.']:
    wn.append([t])
wn['A1'].font = Font(name=FN, bold=True, size=14, color=NAVY); wn.column_dimensions['A'].width = 110
wb.properties.creator = ''; wb.properties.lastModifiedBy = ''
wb.save(f'{ROOT}/excel/Credit_Risk_Analysis.xlsx')


# ---------- 6. DASHBOARD (index.html) ----------
ORD = {'grade': list('ABCDE'), 'dti_band': ['<20%','20-30%','30-40%','40%+'], 'util_band': ['<30%','30-60%','60-80%','80%+'],
       'income_bracket': ['Low (<4L)','Mid (4-8L)','Upper-mid (8-15L)','High (15L+)']}
def seg_rows(d, col):
    t = d.groupby(col).agg(a=('loan_id','size'), p=('approved','sum'), f=('defaulted','sum'))
    return [[k, int(t.a[k]), int(t.p[k]), int(t.f[k])] for k in ORD.get(col, sorted(t.index)) if k in t.index]
def slice_(d, thr):
    d = d.assign(sla=(d.processing_days > 5).astype(int))
    r = d.groupby('region').agg(n=('loan_id','size'), p=('approved','sum'), f=('defaulted','sum'), days=('processing_days','mean'), sla=('sla','sum'), fr=('fraud_flag','sum'))
    pk_ = d.groupby(['grade','dti_band']).agg(n=('loan_id','size'), p=('approved','sum'), f=('defaulted','sum')).reset_index()
    pk_ = pk_[pk_.p >= thr].assign(rt=lambda x: x.f / x.p).sort_values('rt', ascending=False).head(6)
    return dict(n=len(d), appr=int(d.approved.sum()), dflt=int(d.defaulted.sum()),
        disb=float(d.loc[d.approved == 1, 'loan_amount'].sum() / 1e7), loss=float(d.loc[d.defaulted == 1, 'loan_amount'].sum() * .6 / 1e7),
        fraud=int(d.fraud_flag.sum()), sla=int(d.sla.sum()), days=float(d.processing_days.mean()),
        grade=seg_rows(d, 'grade'), dti=seg_rows(d, 'dti_band'), util=seg_rows(d, 'util_band'), inc=seg_rows(d, 'income_bracket'), purpose=seg_rows(d, 'purpose'),
        region=[[k, int(x.n), int(x.p), int(x.f), round(float(x.days), 2), int(x.sla), int(x.fr)] for k, x in r.iterrows()],
        month=[[k, int(v)] for k, v in d.groupby('app_month').size().items()],
        pockets=[[f'Grade {x.grade} | DTI {x.dti_band}', int(x.n), int(x.p), int(x.f)] for x in pk_.itertuples()])
DATA = {'All': slice_(clean, 150)}
for y in ['2023', '2024', '2025']: DATA[y] = slice_(clean[clean.app_month.str[:4] == y], 60)
assert DATA['All']['grade'][-1][3] == int(R['grade'].defaults.iloc[-1]) and DATA['All']['n'] == n_clean

CSS = """:root{--bg:#eef1f5;--cd:#fff;--tx:#1b2433;--mut:#66728a;--bd:#d9dfe9;--nv:#12263f;--acc:#1d4f91;--bad:#b83227;--mid:#d29a2e;--good:#3f7f7a;--soft:#e6ecf5}
@media (prefers-color-scheme:dark){:root{--bg:#0e1420;--cd:#161f30;--tx:#e6ebf4;--mut:#93a0b8;--bd:#263250;--nv:#0a1220;--acc:#6ea2ff;--soft:#1e2a44;--good:#5fb3ab}}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:64px}
body{margin:0;background:var(--bg);color:var(--tx);font:14px/1.5 -apple-system,"Segoe UI",Roboto,Arial,sans-serif;font-variant-numeric:tabular-nums}
header{background:var(--nv);color:#fff;position:sticky;top:0;z-index:5;border-bottom:3px solid var(--bad)}
.hb{max-width:1180px;margin:auto;padding:10px 20px;display:flex;flex-wrap:wrap;gap:10px 24px;align-items:center;justify-content:space-between}
.hb h1{font-size:17px;margin:0;font-weight:650}.hb small{display:block;color:#aebbd3;font-size:12px}
nav{display:flex;gap:16px;flex-wrap:wrap;align-items:center}nav a{color:#cfd9ea;text-decoration:none;font-size:13px}nav a:hover{color:#fff}
select{background:#1d3557;color:#fff;border:1px solid #38527c;border-radius:6px;padding:5px 8px;font:inherit}
main{max-width:1180px;margin:auto;padding:18px 20px 30px}
h2{font-size:12px;letter-spacing:.09em;text-transform:uppercase;color:var(--mut);margin:26px 0 10px;padding-bottom:6px;border-bottom:1px solid var(--bd)}
.brief{background:var(--soft);border-left:4px solid var(--acc);padding:10px 14px;border-radius:4px;margin-bottom:14px}
.kp{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px}
.k{background:var(--cd);border:1px solid var(--bd);border-top:3px solid var(--acc);border-radius:4px;padding:12px 14px}
.k.r{border-top-color:var(--bad)}.k small{display:block;color:var(--mut);font-size:11px;text-transform:uppercase;letter-spacing:.05em}
.k b{font-size:24px;font-weight:650}.k span{display:block;color:var(--mut);font-size:12px}
.g{display:grid;grid-template-columns:repeat(auto-fit,minmax(270px,1fr));gap:12px}.g2{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:12px;margin-top:12px}
.c{background:var(--cd);border:1px solid var(--bd);border-radius:4px;padding:12px 14px;min-width:0}.c h3{margin:0;font-size:14px}.c p{margin:2px 0 6px;color:var(--mut);font-size:12px}
svg{width:100%;height:auto;display:block}svg text{fill:var(--mut);font-size:10px}svg .vl{fill:var(--tx);font-weight:600}
.ax,.gr{stroke:var(--bd)}.avg{stroke:var(--mut);stroke-dasharray:4 3}.area{fill:var(--acc);opacity:.12}.ln{fill:none;stroke:var(--acc);stroke-width:2}.pt{fill:var(--acc)}
.sc{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:13px}th{color:var(--mut);font-weight:600;font-size:11px;text-transform:uppercase;letter-spacing:.04em;text-align:right;padding:6px 8px;border-bottom:2px solid var(--bd)}
td{padding:7px 8px;border-bottom:1px solid var(--bd);text-align:right}th:first-child,td:first-child{text-align:left}
.hm{font-weight:600;color:#fff;border-radius:3px;padding:2px 6px;display:inline-block;min-width:52px;text-align:center}
.hr{display:grid;grid-template-columns:120px 1fr 48px;gap:8px;align-items:center;margin:6px 0;font-size:12px}.hr div{background:var(--soft);height:10px;border-radius:2px}.hr i{display:block;height:100%;background:var(--acc);border-radius:2px}.hr b{text-align:right}
.rc{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:12px}.rc .c{border-left:4px solid var(--good)}.rc h3{margin-bottom:6px}.rc em{font-style:normal;color:var(--mut);font-size:12px;display:block;margin-top:6px}
footer{max-width:1180px;margin:auto;padding:0 20px 28px;color:var(--mut);font-size:12px}"""

JS = r"""
const D=__DATA__,MON=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'];
const $=s=>document.querySelector(s),fmt=n=>Math.round(n).toLocaleString('en-IN'),pc=(a,b)=>b?100*a/b:0,short=s=>s.replace(/ \(.*\)/,'');
const tone=(v,avg)=>v>avg*1.3?'var(--bad)':v<avg*0.8?'var(--good)':'var(--mid)';
function cols(rows,avg,label){const v=rows.map(r=>pc(r[3],r[2])),W=300,H=180,pl=8,pb=26,pt=18,mx=Math.max(...v,avg)*1.18,bw=(W-pl-8)/rows.length,y=t=>pt+(H-pt-pb)*(1-t/mx);
let s='<svg viewBox="0 0 '+W+' '+H+'" role="img" aria-label="'+label+'"><line class="ax" x1="'+pl+'" x2="'+(W-8)+'" y1="'+y(0)+'" y2="'+y(0)+'"/><line class="avg" x1="'+pl+'" x2="'+(W-8)+'" y1="'+y(avg)+'" y2="'+y(avg)+'"/>';
rows.forEach((r,i)=>{const x=pl+i*bw+bw*.16,w=bw*.68;s+='<rect x="'+x+'" y="'+y(v[i])+'" width="'+w+'" height="'+(y(0)-y(v[i]))+'" fill="'+tone(v[i],avg)+'" rx="2"><title>'+short(r[0])+': '+v[i].toFixed(2)+'% default ('+fmt(r[3])+' of '+fmt(r[2])+' approved)</title></rect><text class="vl" x="'+(x+w/2)+'" y="'+(y(v[i])-4)+'" text-anchor="middle">'+v[i].toFixed(1)+'%</text><text x="'+(x+w/2)+'" y="'+(H-9)+'" text-anchor="middle">'+short(r[0])+'</text>'});
return s+'</svg>'}
function card(t,sub,body){return '<div class="c"><h3>'+t+'</h3><p>'+sub+'</p>'+body+'</div>'}
function trend(m){const W=720,H=230,pl=44,pr=10,pt=12,pb=26,n=m.length,mx=Math.max(...m.map(r=>r[1]))*1.12,X=i=>pl+i*(W-pl-pr)/Math.max(n-1,1),Y=v=>pt+(H-pt-pb)*(1-v/mx);
let s='<svg viewBox="0 0 '+W+' '+H+'">';for(let k=0;k<=4;k++){const v=mx*k/4;s+='<line class="gr" x1="'+pl+'" x2="'+(W-pr)+'" y1="'+Y(v)+'" y2="'+Y(v)+'"/><text x="'+(pl-6)+'" y="'+(Y(v)+3)+'" text-anchor="end">'+fmt(v)+'</text>'}
const p=m.map((r,i)=>X(i).toFixed(1)+','+Y(r[1]).toFixed(1)).join(' ');
s+='<polygon class="area" points="'+pl+','+Y(0)+' '+p+' '+X(n-1)+','+Y(0)+'"/><polyline class="ln" points="'+p+'"/>';
m.forEach((r,i)=>{s+='<circle class="pt" cx="'+X(i)+'" cy="'+Y(r[1])+'" r="3"><title>'+MON[+r[0].slice(5)-1]+' '+r[0].slice(0,4)+': '+fmt(r[1])+' applications</title></circle>';if(i%3===0)s+='<text x="'+X(i)+'" y="'+(H-8)+'" text-anchor="middle">'+MON[+r[0].slice(5)-1]+' '+r[0].slice(2,4)+'</text>'});return s+'</svg>'}
function season(m){const yrs=new Set(m.map(r=>r[0].slice(0,4))).size,t=Array(12).fill(0);m.forEach(r=>t[+r[0].slice(5)-1]+=r[1]/yrs);
const W=340,H=180,pb=24,pt=16,mx=Math.max(...t)*1.15,bw=(W-10)/12,top=[...t].sort((a,b)=>b-a).slice(0,3),y=v=>pt+(H-pt-pb)*(1-v/mx);let s='<svg viewBox="0 0 '+W+' '+H+'">';
t.forEach((v,i)=>{const x=5+i*bw+bw*.14,w=bw*.72;s+='<rect x="'+x+'" y="'+y(v)+'" width="'+w+'" height="'+(y(0)-y(v))+'" rx="2" fill="'+(top.includes(v)?'var(--acc)':'var(--bd)')+'"><title>'+MON[i]+': avg '+fmt(v)+' applications</title></rect><text x="'+(x+w/2)+'" y="'+(H-8)+'" text-anchor="middle">'+MON[i][0]+'</text>'});return s+'</svg>'}
function heat(v,avg){return '<span class="hm" style="background:'+tone(v,avg)+'">'+v.toFixed(1)+'%</span>'}
function render(k){const d=D[k],dr=pc(d.dflt,d.appr),g=d.grade,gA=pc(g[0][3],g[0][2]),gE=pc(g[g.length-1][3],g[g.length-1][2]);
const reg=d.region.map(r=>({n:r[0],a:r[1],ap:pc(r[2],r[1]),dr:pc(r[3],r[2]),dy:r[4],sl:pc(r[5],r[1]),fr:pc(r[6],r[1])})).sort((a,b)=>b.dr-a.dr);
const sel=k==='All'?'Jan 2023 - Dec 2025':k;
$('#brief').innerHTML='<b>Summary ('+sel+'):</b> '+fmt(d.n)+' applications, '+pc(d.appr,d.n).toFixed(1)+'% approved, portfolio default rate '+dr.toFixed(1)+'%. Default risk rises from '+gA.toFixed(1)+'% (Grade A) to '+gE.toFixed(1)+'% (Grade E). Highest-risk region: '+reg[0].n+' ('+reg[0].dr.toFixed(1)+'%).';
const K=[['Applications',fmt(d.n),'loan applications received',''],['Approval rate',pc(d.appr,d.n).toFixed(1)+'%',fmt(d.appr)+' approved',''],['Default rate',dr.toFixed(1)+'%',fmt(d.dflt)+' defaults on approved loans','r'],
['Disbursed','Rs '+fmt(d.disb)+' Cr','approved loan value',''],['Estimated credit loss','Rs '+fmt(d.loss)+' Cr','assumes 60% loss given default','r'],['SLA breach',pc(d.sla,d.n).toFixed(1)+'%','decisions taking over 5 days','']];
$('#kpis').innerHTML=K.map(x=>'<div class="k '+x[3]+'"><small>'+x[0]+'</small><b>'+x[1]+'</b><span>'+x[2]+'</span></div>').join('');
const av='Dashed line = portfolio average ('+dr.toFixed(1)+'%). Red: over 1.3x average.';
$('#risk').innerHTML=card('Default rate by credit grade','Grade A = best score band. '+av,cols(g,dr,'grade'))+card('Default rate by debt-to-income','Monthly debt payments as share of income. '+av,cols(d.dti,dr,'dti'))
+card('Default rate by credit utilisation','Share of available credit in use. '+av,cols(d.util,dr,'util'))+card('Default rate by income bracket','Low under Rs 4L, Mid 4-8L, Upper-mid 8-15L, High 15L+. '+av,cols(d.inc,dr,'income'));
const mp=Math.max(...d.pockets.map(r=>pc(r[3],r[2])));
$('#pockets').innerHTML=card('Highest-risk borrower segments','Credit grade and DTI combinations with the highest default rate (minimum approved loans applied).','<div class="sc"><table><tr><th>Segment</th><th>Applications</th><th>Approved</th><th>Defaults</th><th>Default rate</th></tr>'+d.pockets.map(r=>'<tr><td>'+r[0]+'</td><td>'+fmt(r[1])+'</td><td>'+fmt(r[2])+'</td><td>'+fmt(r[3])+'</td><td>'+heat(pc(r[3],r[2]),dr)+'</td></tr>').join('')+'</table></div>')
+card('Default rate by loan purpose','Sorted highest to lowest.',d.purpose.map(r=>({l:r[0],v:pc(r[3],r[2])})).sort((a,b)=>b.v-a.v).map(o=>'<div class="hr"><span>'+o.l+'</span><div><i style="width:'+(o.v/Math.max(...d.purpose.map(r=>pc(r[3],r[2])))*100).toFixed(1)+'%"></i></div><b>'+o.v.toFixed(1)+'%</b></div>').join(''));
$('#demand').innerHTML=card('Monthly loan applications','Demand trend. Hover over a point for the exact count.',trend(d.month))+card('Seasonality','Average applications per calendar month. Blue = top three months.',season(d.month));
$('#regions').innerHTML=card('Regional performance','Sorted by default rate. Colour vs portfolio average default rate.','<div class="sc"><table><tr><th>Region</th><th>Applications</th><th>Approval rate</th><th>Default rate</th><th>Avg decision days</th><th>SLA breach</th><th>Fraud flag</th></tr>'
+reg.map(r=>'<tr><td>'+r.n+'</td><td>'+fmt(r.a)+'</td><td>'+r.ap.toFixed(1)+'%</td><td>'+heat(r.dr,dr)+'</td><td>'+r.dy.toFixed(1)+'</td><td>'+r.sl.toFixed(1)+'%</td><td>'+r.fr.toFixed(2)+'%</td></tr>').join('')+'</table></div>')}
$('#yr').onchange=e=>render(e.target.value);render('All');
"""

BODY = """<header><div class="hb"><div><h1>Retail Credit Risk &amp; Loan Demand Analytics</h1><small>Portfolio review, Jan 2023 - Dec 2025 | Simulated lending dataset</small></div>
<nav><a href="#overview">Executive Overview</a><a href="#drivers">Risk Drivers</a><a href="#ops">Demand &amp; Operations</a><a href="#recs">Recommendations</a>
<label><select id="yr" aria-label="Year"><option value="All">All years</option><option>2023</option><option>2024</option><option>2025</option></select></label></nav></div></header>
<main><h2 id="overview">Executive Overview</h2><div class="brief" id="brief"></div><div class="kp" id="kpis"></div>
<h2 id="drivers">Risk Drivers</h2><div class="g" id="risk"></div><div class="g2" id="pockets"></div>
<h2 id="ops">Operational &amp; Demand Insights</h2><div class="g2" id="demand"></div><div class="g2" id="regions" style="grid-template-columns:1fr"></div>
<h2 id="recs">Business Recommendations</h2><div class="rc">__RECS__</div></main>
<footer>Data notes: the dataset is fully simulated for portfolio and analytical demonstration. Loss estimate assumes 60% loss given default. SLA breach = decision took more than 5 days. Recommendations reflect the full 2023-2025 portfolio.</footer>"""

dr_all = 100 * K.defaults / K.approved
recs = [("Tighten underwriting for high-risk pockets", REC[0].replace('Tighten approval rules or pricing for ', 'Apply stricter approval rules or risk-based pricing to '), F[0] + ' ' + F[1]),
        ("Add a utilisation trigger", "Send applicants above 80% credit utilisation for secondary review before approval.", F[2]),
        ("Staff ahead of the Q4 peak", "Increase processing capacity from September and use monthly demand as a staffing forecast.", F[4] + ' ' + F[5])]
RECS = ''.join(f'<div class="c"><h3>{i}. {t}</h3><div>{a}</div><em>Evidence: {e}</em></div>' for i, (t, a, e) in enumerate(recs, 1))
page = ('<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
        '<title>Retail Credit Risk &amp; Loan Demand Analytics</title><style>' + CSS + '</style></head><body>' + BODY.replace('__RECS__', RECS)
        + '<script>' + JS.replace('__DATA__', json.dumps(DATA, separators=(',', ':'))) + '</script></body></html>')
open(f'{ROOT}/index.html', 'w', encoding='utf-8').write(page)

# ---------- 7. README + personal notes ----------
ar = 100 * K.approved / K.applications
rg_hi, rg_lo = rg.segment.iloc[0], rg.segment.iloc[-1]
top_pk = R['risk_pockets'].iloc[0]
open(f'{ROOT}/README.md', 'w', encoding='utf-8').write(f"""# Retail Credit Risk & Loan Demand Analytics

An end-to-end data analytics project analysing a simulated retail lending portfolio to identify credit-risk drivers, loan-demand patterns and operational bottlenecks.

### Dashboard
[Live Interactive Dashboard](https://YOUR-USERNAME.github.io/retail-credit-risk-analytics/)

### Tools
Python | SQL | Excel | HTML/JavaScript dashboard

> Note: The dataset is fully simulated for portfolio and analytical demonstration purposes. No real customer data is used.

![Risk drivers](docs/screenshots/dashboard-risk.png)
![Demand and operations](docs/screenshots/dashboard-demand.png)

## Project Flow
```
Business Problem
      |
Data Generation (Python)
      |
SQL Cleaning & Segmentation
      |
Python Analysis
      |
Excel Workbook + Interactive Dashboard
      |
Business Insights
      |
Recommendations
```

## Business Problem
A retail lender wants to know which borrowers drive defaults, when loan demand peaks, and where decision turnaround (SLA) is slipping, so it can tighten underwriting where needed and plan capacity ahead of peak months.

## Repository Structure
| Path | Contents |
|---|---|
| `index.html` | Interactive dashboard (Executive Overview, Risk Drivers, Demand & Operations, Recommendations) with year filter |
| `excel/Credit_Risk_Analysis.xlsx` | Dashboard tab with KPIs and charts, formula-driven analysis tab, 10,000-row sample |
| `sql/queries.sql` | Cleaning, deduplication, segmentation and analysis queries (SQLite) |
| `src/credit_risk_analysis.py` | Data simulation, SQL execution, Excel and dashboard build |
| `data/loans_clean.csv` | {n_clean:,} cleaned loan applications |
| `docs/screenshots/` | Dashboard screenshots |

## Data Cleaning (SQL)
- Raw records: {n_raw:,}. Removed {dups} duplicate loan IDs and {bad} invalid credit scores. Final dataset: {n_clean:,} records.
- Filled {n_inc:,} missing or outlier incomes and {n_dti:,} missing DTI values using regional averages.
- Created segments: credit grade (A-E), DTI band, credit utilisation band, income bracket and application month.

## Key Findings
Portfolio: {int(K.applications):,} applications, {ar:.1f}% approved, {dr_all:.1f}% default rate, estimated credit loss of Rs {int(K.est_loss_cr):,} Cr.

{chr(10).join('- ' + t for t in F)}
- Highest-risk segment: {top_pk.segment} with a {top_pk.default_pct:.1f}% default rate.

## Recommendations
{chr(10).join(f'{i}. ' + t for i, t in enumerate(REC, 1))}

## Assumptions and Limitations
- Data is simulated; relationships between variables are built into the generator, so results demonstrate the method rather than real market behaviour.
- Loss given default is assumed at 60% of loan value. SLA breach is defined as a decision taking more than 5 days.
- Default is a binary flag; time-to-default and probability-of-default modelling are out of scope.

## Run Locally
```
pip install -r requirements.txt
python src/credit_risk_analysis.py
```
Open `index.html` in a browser to view the dashboard.
""")
open(f'{ROOT}/requirements.txt', 'w').write('pandas\nnumpy\nopenpyxl\n')
open(f'{ROOT}/.gitignore', 'w').write('__pycache__/\n*.pyc\n.DS_Store\n')
open(os.path.join(os.path.dirname(ROOT), 'MY_NOTES_do_not_upload.md'), 'w', encoding='utf-8').write(f"""# Your private notes (do NOT upload this file to GitHub)

## 1. Put it on GitHub and make it live (about 5 minutes, no coding)
1. Unzip `retail-credit-risk-analytics.zip`. You will get a folder with the same name.
2. github.com > New repository. Name: `retail-credit-risk-analytics`. Public. Create (leave README unticked).
3. Click "uploading an existing file". Drag in EVERYTHING INSIDE the folder (index.html, README.md, src, data, sql, excel, docs, requirements.txt). Folders upload with drag and drop. Commit.
4. Repo > Settings > Pages > Source: "Deploy from a branch" > Branch `main`, folder `/ (root)` > Save. After about a minute your dashboard is live at `https://YOUR-USERNAME.github.io/retail-credit-risk-analytics/`.
5. Open README.md on GitHub > pencil icon > replace `YOUR-USERNAME` in the dashboard link with your GitHub username > Commit.
6. Set the repo "About" box (gear icon, top right) website to the live link.

## 2. Resume entry
**Retail Credit Risk & Loan Demand Analytics** | Python, SQL, Excel | GitHub | Live Dashboard
- Built an end-to-end credit risk and loan demand analysis on a simulated {n_clean:,}-record lending portfolio using SQL, Python and Excel, quantifying default rates across credit grade, DTI, utilisation, income and region.
- Developed SQL cleaning and segmentation pipelines to remove duplicates, handle invalid values and missing data, and identify high-risk borrower segments.
- Built an interactive dashboard and Excel workbook tracking approval rate, default rate, loan demand, regional performance and SLA breaches.
- Identified Grade E and high-DTI borrowers as key risk pockets ({top_pk.segment}: {top_pk.default_pct:.1f}% default) and recommended underwriting controls and pre-Q4 staffing based on demand seasonality.

Add "Power BI" to the tools and bullets ONLY after you finish section 3.

## 3. Power BI (I cannot create a .pbix file; this takes about 20 minutes)
1. Power BI Desktop > Get data > Text/CSV > `data/loans_clean.csv` > Load.
2. Modeling > New measure (paste each):
   - `Applications = COUNTROWS(loans_clean)`
   - `Approved = SUM(loans_clean[approved])`
   - `Approval Rate = DIVIDE([Approved], [Applications])`
   - `Defaults = SUM(loans_clean[defaulted])`
   - `Default Rate = DIVIDE([Defaults], [Approved])`
   - `SLA Breach % = DIVIDE(CALCULATE(COUNTROWS(loans_clean), loans_clean[processing_days] > 5), [Applications])`
3. Page 1 "Risk": cards (Applications, Approval Rate, Default Rate); clustered columns of Default Rate by `grade`, `dti_band`, `util_band`, `region`; slicer on `app_month` year.
4. Page 2 "Demand": line chart `app_month` vs Applications; table of region with Approval Rate, Default Rate, SLA Breach %; slicer on `purpose`.
5. Save as `powerbi/Credit_Risk.pbix`, upload it to the repo, take a screenshot into `docs/screenshots/`, and add "Power BI" to README Tools line and resume.

## 4. Interview answers
- Why simulated data? To control data-quality problems and keep the project reproducible; the method transfers directly to real data.
- What is DTI? Debt-to-income: monthly debt payments divided by monthly income. Higher means more repayment stress.
- Most important finding? Grade and DTI together isolate the riskiest pockets, so targeted tightening cuts losses without cutting volume everywhere.
- What would you improve? Model probability of default (logistic regression), add vintage analysis, and test on real data.
- How did you clean the data? SQL CTEs: deduplicate on loan ID, remove impossible scores, impute missing income and DTI with regional averages, treat income above Rs 1 crore as outliers.
""")
print('OK', n_raw, n_clean)
